import hashlib
import io
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook

from app.importers.moneylover import HEADERS, parse_moneylover_xlsx


def workbook_bytes(headers=HEADERS, rows=()):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sổ giao dịch"
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_parses_rows_and_preserves_source_data() -> None:
    data = workbook_bytes(
        rows=[
            (
                "id-1",
                date(2026, 8, 22),
                "Food",
                12.5,
                "USD",
                "Cash",
                "note",
                "",
                "",
                False,
                "member",
            )
        ]
    )
    parsed = parse_moneylover_xlsx(data)
    row = parsed.rows[0]
    assert row.source_row_number == 2
    assert row.source_row_id == "id-1"
    assert row.transaction_date == date(2026, 8, 22)
    assert not isinstance(row.transaction_date, datetime)
    assert row.amount == Decimal("12.5")
    assert row.amount_scaled == 125000
    assert row.raw_payload["Ghi chú"] == "note"
    assert parsed.file_sha256 == hashlib.sha256(data).hexdigest()
    assert parse_moneylover_xlsx(io.BytesIO(data)).file_sha256 == parsed.file_sha256


def test_requires_sheet_and_headers() -> None:
    workbook = Workbook()
    workbook.active.title = "Summary"
    output = io.BytesIO()
    workbook.save(output)
    with pytest.raises(ValueError, match="missing required worksheet"):
        parse_moneylover_xlsx(output.getvalue())

    with pytest.raises(ValueError, match="invalid"):
        parse_moneylover_xlsx(workbook_bytes(headers=("wrong",)))


def test_tolerates_header_whitespace_and_missing_member_column() -> None:
    headers_with_spaces_and_no_member = (
        "Id ",
        " Ngày",
        "Nhóm ",
        "Số tiền ",
        "Đơn vị tiền tệ",
        "Ví ",
        "Ghi chú",
        "Với",
        "Sự kiện",
        "Không tính vào báo cáo ",
    )
    data = workbook_bytes(
        headers=headers_with_spaces_and_no_member,
        rows=[
            ("id-2", "2026-08-25", "Ăn uống", 45000, "VND", "UOB", "Ăn trưa", "", "", False)
        ],
    )
    parsed = parse_moneylover_xlsx(data)
    assert len(parsed.rows) == 1
    assert parsed.rows[0].source_row_id == "id-2"
    assert parsed.rows[0].transaction_date == date(2026, 8, 25)
    assert parsed.rows[0].amount == Decimal(45000)
    assert parsed.rows[0].raw_payload["Ví"] == "UOB"


def test_tolerates_blank_row_before_headers() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sổ giao dịch"
    sheet.append([])  # blank row 1
    sheet.append(list(HEADERS))  # headers on row 2
    sheet.append(["id-3", "2026-08-26", "Đi lại", 20000, "VND", "Cash", "Gửi xe", "", "", False, ""])
    buf = io.BytesIO()
    workbook.save(buf)

    parsed = parse_moneylover_xlsx(buf.getvalue())
    assert len(parsed.rows) == 1
    assert parsed.rows[0].source_row_number == 3
    assert parsed.rows[0].amount == Decimal(20000)
    assert parsed.rows[0].raw_payload["Nhóm"] == "Đi lại"

